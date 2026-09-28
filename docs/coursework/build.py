"""Сборка курсовой работы по ОП.07 «Экономика отрасли» в .docx.

Оформление — по «Методическим рекомендациям по КР ОП.07» (ГОСТ 7.32-2017):
A4, Times New Roman 14, интервал 1,5, выравнивание по ширине, абзац 1,25 см,
поля 30/15/20/20 мм, нумерация страниц по центру снизу (кроме титульного листа),
в таблицах — 12 пт и одинарный интервал.

Запуск:  python docs/coursework/build.py [выходной .docx]
"""

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).parent))
from content import DOC  # noqa: E402

FONT = "Times New Roman"
BODY = Pt(14)
SMALL = Pt(12)
INDENT = Cm(1.25)
RIGHT_TAB = Cm(16.5)  # ширина полосы набора при полях 30/15 мм


# --- низкоуровневые помощники -------------------------------------------------


def _field(paragraph, instr: str) -> None:
    """Вставить поле Word (PAGE, TOC …) в абзац."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr_el, sep, end):
        run._r.append(el)


def _update_fields_on_open(document) -> None:
    """Word/LibreOffice предложат обновить СОДЕРЖАНИЕ при открытии файла."""
    settings = document.settings.element
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    settings.append(el)


def _no_spacing(paragraph, line: float = 1.5) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = line


def _east_asia(style) -> None:
    """Тот же шрифт для восточноазиатского набора — иначе Word подставит свой."""
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), FONT)
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)


def _setup_styles(document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = BODY
    _east_asia(normal)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = INDENT
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name in ("Heading 1", "Heading 2"):
        st = document.styles[name]
        st.font.name = FONT
        st.font.size = BODY
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.all_caps = False
        _east_asia(st)
        hpf = st.paragraph_format
        hpf.line_spacing = 1.5
        hpf.space_before = Pt(0)
        hpf.space_after = Pt(0)
        hpf.keep_with_next = True

    h1 = document.styles["Heading 1"].paragraph_format
    h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.first_line_indent = Cm(0)
    h2 = document.styles["Heading 2"].paragraph_format
    h2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.first_line_indent = INDENT


def _setup_page(section) -> None:
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(30)
    section.right_margin = Mm(15)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)


# --- операции содержимого -----------------------------------------------------


class Builder:
    def __init__(self) -> None:
        self.doc = Document()
        _setup_styles(self.doc)
        section = self.doc.sections[0]
        _setup_page(section)
        section.different_first_page_header_footer = True  # титул без номера
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.first_line_indent = Cm(0)
        _field(footer, "PAGE")
        for run in footer.runs:
            run.font.name = FONT
            run.font.size = BODY
        _update_fields_on_open(self.doc)

    # текст ------------------------------------------------------------------

    def p(self, text: str, *, indent: bool = True, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
          bold: bool = False, line: float = 1.5, spacing_after: int = 0):
        par = self.doc.add_paragraph()
        _no_spacing(par, line)
        par.paragraph_format.space_after = Pt(spacing_after)
        par.paragraph_format.first_line_indent = INDENT if indent else Cm(0)
        par.alignment = align
        run = par.add_run(text)
        run.bold = bold
        return par

    def center(self, text: str, *, bold: bool = False, line: float = 1.5):
        return self.p(text, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER,
                      bold=bold, line=line)

    def right(self, text: str, *, line: float = 1.0):
        return self.p(text, indent=False, align=WD_ALIGN_PARAGRAPH.RIGHT, line=line)

    def h1(self, text: str, *, new_page: bool = True):
        if new_page:
            self.page_break()
        par = self.doc.add_paragraph(style="Heading 1")
        _no_spacing(par)
        par.paragraph_format.first_line_indent = Cm(0)
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.add_run(text.upper()).bold = True
        return par

    def h2(self, text: str):
        par = self.doc.add_paragraph(style="Heading 2")
        _no_spacing(par)
        par.paragraph_format.first_line_indent = INDENT
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        par.add_run(text).bold = True
        return par

    def dash_list(self, items):
        for item in items:
            self.p("– " + item)

    def num_list(self, items):
        for i, item in enumerate(items, 1):
            self.p(f"{i}) {item}")

    # формулы ----------------------------------------------------------------

    def formula(self, body: str, number: int):
        par = self.doc.add_paragraph()
        _no_spacing(par)
        par.paragraph_format.first_line_indent = Cm(0)
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        tabs = par.paragraph_format.tab_stops
        tabs.add_tab_stop(Cm(8.25), WD_TAB_ALIGNMENT.CENTER)
        tabs.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
        par.add_run(f"\t{body}\t({number})")
        return par

    def where(self, lines):
        """Расшифровка обозначений под формулой: «где» без абзацного отступа."""
        first, *rest = lines
        self.p("где " + first, indent=False)
        for line in rest:
            self.p(line, indent=False)

    # таблицы ----------------------------------------------------------------

    def caption(self, text: str):
        par = self.p(text, indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.0)
        par.paragraph_format.space_before = Pt(6)
        return par

    def table(self, rows, widths=None, header_bold: bool = True):
        cols = len(rows[0])
        table = self.doc.add_table(rows=0, cols=cols)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        for r, row in enumerate(rows):
            cells = table.add_row().cells
            for c, value in enumerate(row):
                cell = cells[c]
                cell.text = ""
                par = cell.paragraphs[0]
                _no_spacing(par, line=1.0)
                par.paragraph_format.first_line_indent = Cm(0)
                par.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER if r == 0 else WD_ALIGN_PARAGRAPH.LEFT
                )
                run = par.add_run(str(value))
                run.font.name = FONT
                run.font.size = SMALL
                run.bold = header_bold and r == 0
        if widths:
            for row in table.rows:
                for cell, width in zip(row.cells, widths):
                    cell.width = Cm(width)
        after = self.doc.add_paragraph()
        _no_spacing(after)
        after.paragraph_format.first_line_indent = Cm(0)
        return table

    def page_break(self):
        par = self.doc.add_paragraph()
        _no_spacing(par)
        par.paragraph_format.first_line_indent = Cm(0)
        par.add_run().add_break(WD_BREAK.PAGE)
        return par

    def blank(self, n: int = 1):
        for _ in range(n):
            par = self.doc.add_paragraph()
            _no_spacing(par)
            par.paragraph_format.first_line_indent = Cm(0)

    def toc(self):
        par = self.doc.add_paragraph()
        _no_spacing(par)
        par.paragraph_format.first_line_indent = Cm(0)
        _field(par, 'TOC \\o "1-2" \\h \\z \\u')

    def signature_table(self, rows):
        table = self.doc.add_table(rows=0, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row in rows:
            cells = table.add_row().cells
            for c, value in enumerate(row):
                par = cells[c].paragraphs[0]
                _no_spacing(par, line=1.0)
                par.paragraph_format.first_line_indent = Cm(0)
                run = par.add_run(value)
                run.font.name = FONT
                run.font.size = BODY
        for row in table.rows:
            for cell, width in zip(row.cells, (7.0, 4.5, 5.0)):
                cell.width = Cm(width)
        return table

    def save(self, path: Path):
        self.doc.save(str(path))


OPS = {
    "h1": lambda b, a: b.h1(*a[:1], **(a[1] if len(a) > 1 else {})),
    "h2": lambda b, a: b.h2(a[0]),
    "p": lambda b, a: b.p(a[0]),
    "pn": lambda b, a: b.p(a[0], indent=False),
    "center": lambda b, a: b.center(a[0], bold=(len(a) > 1 and a[1])),
    "right": lambda b, a: b.right(a[0]),
    "dash": lambda b, a: b.dash_list(a[0]),
    "num": lambda b, a: b.num_list(a[0]),
    "f": lambda b, a: b.formula(a[0], a[1]),
    "where": lambda b, a: b.where(a[0]),
    "cap": lambda b, a: b.caption(a[0]),
    "tbl": lambda b, a: b.table(a[0], a[1] if len(a) > 1 else None),
    "toc": lambda b, a: b.toc(),
    "break": lambda b, a: b.page_break(),
    "blank": lambda b, a: b.blank(a[0] if a else 1),
    "sign": lambda b, a: b.signature_table(a[0]),
}


def main() -> None:
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else (
        Path(__file__).resolve().parents[2]
        / "Курсовая 1 ОП.07 Экономика отрасли (mosres, Android-клиент).docx"
    )
    b = Builder()
    for op, *args in DOC:
        OPS[op](b, args)
    b.save(out)
    print(f"готово: {out}")


if __name__ == "__main__":
    main()
